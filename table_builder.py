import openpyxl
from openpyxl.styles import PatternFill, Alignment, Border, Side

def combine_sequenced_lectures(week):
    for day in week:
        for hour in week[day]:
            for l, lec in enumerate(week[day][hour]):
                week[day][hour][l]["count"] = 1

    for day in week:
        reversed_hours = list(week[day].keys())[::-1]
        for h, hour in enumerate(reversed_hours[:-1]):
            for l in range(len(week[day][hour])-1, -1, -1):
                lec = week[day][hour][l]
                for l2, lec2 in enumerate(week[day][reversed_hours[h+1]]):
                    if lec["code"] == lec2["code"] and lec["professorNumber"] == lec2["professorNumber"]:
                        week[day][reversed_hours[h+1]][l2]["count"] += week[day][hour][l]["count"]
                        week[day][hour].pop(l)
                        break
    
    return week

def tableize_combined_week_by_year(combined_week, remove_unnecessary_nones = False):
    cols_availibility_counter = {}
    table = {
        "rows": {},
        "cols": {},
    }
    global_row_pointer = 0

    for day in combined_week:
        table["rows"][day] = {}

        for hour in combined_week[day]:
            table["rows"][day][hour] = {}

            if not remove_unnecessary_nones:
                for year in cols_availibility_counter:
                    for i, item in enumerate(cols_availibility_counter[year]):
                        if item != 0:
                            table["cols"][year][i].append(None)

            lectures = combined_week[day][hour][::-1]

            for l in range(len(lectures)-1, -1, -1):
                lec = lectures[l]

                if lec["year"] not in cols_availibility_counter:
                    cols_availibility_counter[lec["year"]] = []
                    table["cols"][lec["year"]] = []

                for year in cols_availibility_counter:

                    if lec["year"] != year:
                        continue

                    for i, item in enumerate(cols_availibility_counter[year]):
                        if item == 0:
                            table["cols"][year][i].append(lec)
                            cols_availibility_counter[year][i] = lec["count"]
                            lectures.pop(l)
                            break
                    else:
                        table["cols"][year].append([{} for _ in range(global_row_pointer)] + [lec])
                        cols_availibility_counter[year].append(lec["count"])


            for year in cols_availibility_counter:
                for i, item in enumerate(cols_availibility_counter[year]):
                    if item == 0:
                        table["cols"][year][i].append({})
                    else:
                        cols_availibility_counter[year][i] -= 1

            global_row_pointer += 1

    table["cols"] = dict(sorted(table["cols"].items())) # to sort keys

    return table

def build_week_html_content(tableized_week):
    style = """<style>
    th, tr {
        border: 1px solid black;
    }
    td, th {
        padding: 5px;
    }
    .bordered {
        border: 1px solid black;
        text-align: center;
        vertical-align: middle;
    }
    .vertical_writing {
        transform: rotate(-90deg);
    }
    table {
        border-collapse: collapse;
    }
    .header_cell {
        /* background-color: rgb(255, 89, 89); */
        background-color: aqua;
    }
    .lec_cell {
        background-color: rgb(224, 224, 224);
    }
    </style>"""

    table_content = style + "<table>"

    header_content = """<tr>
    <th class="bordered header_cell">Day</th>
    <th class="bordered header_cell">Time</th>""" + "".join(
        f"""
        <th class="bordered header_cell" colspan="{len(tableized_week["cols"][year])}">Year {year}</th>
        """
        for year in tableized_week["cols"]
    ) + """
</tr>"""
    

    global_table_row_pointer = 0

    for day in tableized_week["rows"]:
        table_content += header_content
        
        days_content = f"""<td class="bordered vertical_writing" rowspan="{len(tableized_week["rows"][day])+1}">{day}</td>"""
        
        rows_content = []
        for hour in tableized_week["rows"][day]:
            rows_content.append(f"""<tr><td class="bordered">{hour}:00 ~ {hour}:50</td>""")
        
        for year in tableized_week["cols"]:
            for col in tableized_week["cols"][year]:
                for l, lec in enumerate(col[global_table_row_pointer : global_table_row_pointer + len(rows_content)]):
                    if lec is None:
                        continue

                    if lec == {}:
                        rows_content[l] += "<td></td>"
                        continue

                    span = lec["count"]

                    rows_content[l] += f"""<td rowspan="{span}" class="bordered lec_cell">{lec["name"]}<br>{lec["professorName"]}</td>"""
        
        global_table_row_pointer += len(rows_content)

        for i in range(len(rows_content)):
            rows_content[i] += "</tr>"
        
        days_content += "".join(rows_content)

        table_content += days_content

    table_content += "</table>"

    return table_content

def auto_adjust_column_width(ws):
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter  # Get column letter (A, B, C, etc.)
        
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        
        ws.column_dimensions[col_letter].width = max_length

def auto_adjust_row_height(ws):
    for row in ws.iter_rows():
        max_height = 15  # Default row height
        for cell in row:
            if cell.value:
                lines = str(cell.value).count("\n") + 1  # Count line breaks
                max_height = max(max_height, lines * 15)  # Adjust for multiline text
        
        ws.row_dimensions[row[0].row].height = max_height

def get_max_dimentions(tableized_week):
    rows_num = 0
    for day in tableized_week["rows"]:
        rows_num += len(tableized_week["rows"][day]) + 1 # fore headers

    cols_num = 2
    for year in tableized_week["cols"]:
        years_columns_num = len(tableized_week["cols"][year])
        cols_num += years_columns_num
    
    return rows_num, cols_num

def build_week_excel_file(tableized_week):
    header_fill = PatternFill(start_color="00FFFF", end_color="00FFFF", fill_type="solid")
    cell_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")  # Center alignment
    day_alignment = Alignment(textRotation=90, horizontal="center", vertical="center")  # Center alignment
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    wb = openpyxl.Workbook()
    ws = wb.active

    def add_header(ws, row, col):
        cell = ws.cell(row=1+row, column=1+col, value=f"Day")
        cell.fill = header_fill
        cell.alignment = alignment
        cell.border = border

        cell = ws.cell(row=1+row, column=1+col+1, value=f"Time")
        cell.fill = header_fill
        cell.alignment = alignment
        cell.border = border

        col += 2

        for year in tableized_week["cols"]:
            years_columns_num = len(tableized_week["cols"][year])
            
            cell = ws.cell(row=1+row, column=1+col, value=f"Year {year}")
            cell.fill = header_fill
            cell.alignment = alignment
            cell.border = border

            ws.merge_cells(
                start_row=1+row,
                start_column=1+col,
                end_row=1+row,
                end_column=1+col+years_columns_num -1
            )

            col += years_columns_num

    global_row_pointer = 0
    global_table_row_pointer = 0
    global_col_pointer = 0


    for day in tableized_week["rows"]:
        add_header(ws, global_row_pointer, global_col_pointer)
        global_row_pointer += 1

        cell = ws.cell(row=1+global_row_pointer, column=1+global_col_pointer, value=day)
        # cell.fill = cell_fill
        cell.alignment = day_alignment
        cell.border = border

        ws.merge_cells(
            start_row=1+global_row_pointer,
            start_column=1+global_col_pointer,
            end_row=1+global_row_pointer+len(tableized_week["rows"][day]) -1,
            end_column=1+global_col_pointer
        )

        local_row_pointer = global_row_pointer
        for hour in tableized_week["rows"][day]:
            cell = ws.cell(row=1+local_row_pointer, column=1+global_col_pointer+1, value=f"{hour}:00 ~ {hour}:50")
            # cell.fill = cell_fill
            cell.alignment = alignment
            cell.border = border
            local_row_pointer += 1
        
        hours_num = len(tableized_week["rows"][day])

        local_col_pointer = global_col_pointer
        for year in tableized_week["cols"]:
            # to_skip = 0
            for col in tableized_week["cols"][year]:
                local_row_pointer = global_row_pointer
                for lec in col[global_table_row_pointer : global_table_row_pointer + hours_num]:
                    local_row_pointer += 1


                    # if to_skip > 0:
                    #     to_skip -= 1
                    #     continue

                    if lec is None:
                        continue

                    if lec == {}:
                        continue

                    to_skip = lec["count"] -1
                    
                    cell = ws.cell(
                        row=local_row_pointer,
                        column=1+local_col_pointer+2,
                        value=lec["name"] + "\n" + lec["professorName"]
                    )
                    cell.fill = cell_fill
                    cell.alignment = alignment
                    cell.border = border
                    ws.merge_cells(
                        start_row=local_row_pointer,
                        start_column=1+local_col_pointer+2,
                        end_row=local_row_pointer+to_skip,
                        end_column=1+local_col_pointer+2
                    )

                    # local_row_pointer += to_skip
                
                local_col_pointer += 1
        
        global_row_pointer += hours_num
        global_table_row_pointer += hours_num

    rows_num, cols_num = get_max_dimentions(tableized_week)

    for i in range(rows_num):
        cell = ws.cell(row=1+i, column=cols_num)
        b = cell.border
        cell.border = Border(left=b.left, top=b.top, right=Side(style="thin"), bottom=b.bottom)

    for i in range(cols_num):
        cell = ws.cell(row=rows_num, column=1+i)
        b = cell.border
        cell.border = Border(left=b.left, top=b.top, right=b.right, bottom=Side(style="thin"))

    auto_adjust_column_width(ws)
    # auto_adjust_row_height(ws)

    # wb.save("merged_cells.xlsx")
    return wb

def build_time_table_html_content(times_list):
    def is_in_list(day, hour):
        item = next((item for item in times_list if item.hour == hour and item.day == day), None)

        if item is None:
            return False

        return True


    days = [
        "Monday",
        "Tuesday",
        "Wednesday",
        "Thursday",
        "Friday",
    ]

    hours = [
        8, 9, 10, 11, 13, 14, 15, 16, 17
    ]

    style = """<style>
table {
    border-collapse: collapse;
}
td, tr, th {
    border: 1px solid black;
    padding: 5px;
    text-align: center;
    vertical-align: middle;
}
tr th {
    background-color: aqua;
}
button {
    margin: 2px;
}
.selected {
    background-color: green;
}</style>"""

    table_content = style + "<table id='schedule-table'>"

    header_content = """<thead>
<tr>
    <th>Time</th>""" + "".join(f"""
    <th value="{day}">{day}</th>
    """
    for day in days
) + """
</tr>
</thead>"""

    body_content = """<tbody>""" + "".join(
"<tr>" + f"""<td value="{hour}">{hour}:00 ~ {hour}:50</td>""" + "".join(
        f"""<td class="{is_in_list(day, hour) * "selected"}" onclick="this.classList.toggle('selected');"></td>"""
        for day in days
    ) + "</tr>"
    for hour in hours
) + """
</tbody>"""
        # if hour == 12
        #         <td class="break"></td>
        #     else
        # f"""<td class="{is_in_list(day, hour) * "selected"}" onclick="this.classList.toggle('selected');"></td>"""
    
    foot_content = """<tfoot>
    <tr>
        <td>
            <button onclick="saveFreeTime()">Save</button>
        </td>""" + "".join(
        f"""<td>
    <button onclick="eraseDay({d+1})">Erase</button>
    <button onclick="toggleDay({d+1})">Toggle</button>
</td>"""
        for d in range(len(days))
    ) + """</tr>
</tfoot>"""

    table_content += header_content + body_content + foot_content

    table_content += "</table>"

    return table_content